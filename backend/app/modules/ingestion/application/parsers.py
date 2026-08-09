from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from typing import cast
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import JsonValue

from app.modules.pricing.domain.currencies import normalize_currency_code

MAX_SPREADSHEET_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX_BYTES = 50 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000

_SUPPORTED_COLUMNS = {
    "operation",
    "product_sku",
    "product_name",
    "description",
    "brand",
    "category",
    "collection",
    "variant_sku",
    "size",
    "color",
    "attributes",
    "price",
    "currency",
    "quantity_on_hand",
    "image_references",
    "metadata",
}
_REQUIRED_COLUMNS = {
    "product_sku",
    "product_name",
    "category",
    "variant_sku",
    "price",
    "currency",
    "quantity_on_hand",
}
_SKU = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,63}$")


@dataclass(frozen=True, slots=True)
class ParsedError:
    row_number: int | None
    field: str | None
    code: str
    message: str


@dataclass(frozen=True, slots=True)
class ParsedRow:
    row_number: int
    fingerprint: str
    data: dict[str, JsonValue]


@dataclass(frozen=True, slots=True)
class ParsedSpreadsheet:
    rows: list[ParsedRow]
    errors: list[ParsedError]
    checksum_sha256: str


class SpreadsheetParser:
    def parse(self, filename: str, content_type: str, data: bytes) -> ParsedSpreadsheet:
        safe_name = safe_filename(filename)
        if not data or len(data) > MAX_SPREADSHEET_BYTES:
            raise ValueError("Spreadsheet must contain at most 5 MiB of data.")
        suffix = PurePosixPath(safe_name).suffix.casefold()
        if suffix == ".csv":
            records = self._csv(data)
        elif suffix == ".xlsx":
            records = self._xlsx(data)
        else:
            raise ValueError("Only CSV and XLSX spreadsheets are supported.")
        if content_type and content_type not in {
            "text/csv",
            "application/csv",
            "application/vnd.ms-excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
        }:
            raise ValueError("Spreadsheet content type is not supported.")
        return self._normalize(records, hashlib.sha256(data).hexdigest())

    @staticmethod
    def _csv(data: bytes) -> list[tuple[int, dict[str, object]]]:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("CSV files must use UTF-8 encoding.") from error
        reader = csv.DictReader(StringIO(text), strict=True)
        if reader.fieldnames is None:
            raise ValueError("Spreadsheet header is missing.")
        headers = [_header(value) for value in reader.fieldnames]
        if len(headers) != len(set(headers)):
            raise ValueError("Spreadsheet headers must be unique.")
        return [
            (number, dict(zip(headers, row.values(), strict=True)))
            for number, row in enumerate(reader, start=2)
        ]

    @staticmethod
    def _xlsx(data: bytes) -> list[tuple[int, dict[str, object]]]:
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > (
                    MAX_UNCOMPRESSED_XLSX_BYTES
                ):
                    raise ValueError("Expanded XLSX content is too large.")
            workbook = load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
        except (BadZipFile, KeyError, OSError) as error:
            raise ValueError("XLSX file could not be parsed.") from error
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("XLSX workbook does not contain a worksheet.")
            iterator = sheet.iter_rows()
            header_cells = next(iterator, None)
            if header_cells is None:
                raise ValueError("Spreadsheet header is missing.")
            headers = [_header(cell.value) for cell in header_cells]
            if len(headers) != len(set(headers)):
                raise ValueError("Spreadsheet headers must be unique.")
            records: list[tuple[int, dict[str, object]]] = []
            for number, cells in enumerate(iterator, start=2):
                if any(cell.data_type == "f" for cell in cells):
                    records.append((number, {"__formula__": "formula"}))
                    continue
                values = [cell.value for cell in cells]
                if len(values) < len(headers):
                    values.extend([None] * (len(headers) - len(values)))
                records.append(
                    (number, dict(zip(headers, values[: len(headers)], strict=True)))
                )
            return records
        finally:
            workbook.close()

    def _normalize(
        self, records: list[tuple[int, dict[str, object]]], checksum: str
    ) -> ParsedSpreadsheet:
        if len(records) > MAX_IMPORT_ROWS:
            raise ValueError("Spreadsheet exceeds the 10,000 row limit.")
        rows: list[ParsedRow] = []
        errors: list[ParsedError] = []
        seen_variants: set[str] = set()
        product_shapes: dict[str, tuple[str, str | None, str | None, str]] = {}
        for number, record in records:
            if not any(value not in {None, ""} for value in record.values()):
                continue
            row, row_errors = self._row(number, record)
            if row is not None:
                variant_sku = str(row["variant_sku"])
                if variant_sku in seen_variants:
                    row_errors.append(
                        ParsedError(
                            number,
                            "variant_sku",
                            "duplicate_variant_sku",
                            "Variant SKU is duplicated in this spreadsheet.",
                        )
                    )
                seen_variants.add(variant_sku)
                product_sku = str(row["product_sku"])
                shape = (
                    str(row["product_name"]),
                    _optional_text(row.get("description")),
                    _optional_text(row.get("brand")),
                    str(row["category"]),
                )
                previous = product_shapes.setdefault(product_sku, shape)
                if previous != shape:
                    row_errors.append(
                        ParsedError(
                            number,
                            "product_sku",
                            "inconsistent_product",
                            "Rows for one Product SKU must use consistent "
                            "product data.",
                        )
                    )
            if row_errors:
                errors.extend(row_errors)
                continue
            assert row is not None
            encoded = json.dumps(row, sort_keys=True, separators=(",", ":"))
            rows.append(
                ParsedRow(
                    row_number=number,
                    fingerprint=hashlib.sha256(encoded.encode()).hexdigest(),
                    data=row,
                )
            )
        if not rows and not errors:
            errors.append(
                ParsedError(None, None, "empty_spreadsheet", "No data rows found.")
            )
        return ParsedSpreadsheet(rows=rows, errors=errors, checksum_sha256=checksum)

    def _row(
        self, number: int, record: dict[str, object]
    ) -> tuple[dict[str, JsonValue] | None, list[ParsedError]]:
        errors: list[ParsedError] = []
        if "__formula__" in record:
            return None, [
                ParsedError(
                    number,
                    None,
                    "formula_not_allowed",
                    "Spreadsheet formulas and macros are not accepted.",
                )
            ]
        unknown = set(record) - _SUPPORTED_COLUMNS
        if unknown:
            errors.append(
                ParsedError(
                    number,
                    None,
                    "unknown_columns",
                    f"Unsupported columns: {', '.join(sorted(unknown))}.",
                )
            )
        values: dict[str, str] = {}
        for key, value in record.items():
            try:
                values[key] = _cell(value)
            except ValueError:
                values[key] = ""
                errors.append(
                    ParsedError(
                        number,
                        key,
                        "formula_not_allowed",
                        "Formula-like spreadsheet values are not accepted.",
                    )
                )
        for field in _REQUIRED_COLUMNS:
            if not values.get(field):
                errors.append(
                    ParsedError(number, field, "required", f"{field} is required.")
                )
        for field in ("product_sku", "variant_sku"):
            value = values.get(field)
            if value and not _SKU.fullmatch(value):
                errors.append(
                    ParsedError(number, field, "invalid_sku", f"{field} is invalid.")
                )
        operation = (values.get("operation") or "auto").casefold()
        if operation not in {"auto", "create", "update"}:
            errors.append(
                ParsedError(
                    number,
                    "operation",
                    "invalid_operation",
                    "Operation must be auto, create, or update.",
                )
            )
        try:
            price = Decimal(values.get("price") or "")
            exponent = cast(int, price.as_tuple().exponent)
            if not price.is_finite() or price < 0 or exponent < -4:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            errors.append(
                ParsedError(number, "price", "invalid_price", "Price is invalid.")
            )
            price = Decimal(0)
        try:
            currency = normalize_currency_code(values.get("currency") or "")
        except ValueError:
            errors.append(
                ParsedError(
                    number, "currency", "invalid_currency", "Currency is invalid."
                )
            )
            currency = "XXX"
        try:
            quantity = int(values.get("quantity_on_hand") or "")
            if quantity < 0 or str(quantity) != (values.get("quantity_on_hand") or ""):
                raise ValueError
        except ValueError:
            errors.append(
                ParsedError(
                    number,
                    "quantity_on_hand",
                    "invalid_quantity",
                    "Quantity must be a non-negative integer.",
                )
            )
            quantity = 0
        attributes, attribute_error = _attributes(values)
        if attribute_error:
            errors.append(
                ParsedError(number, "attributes", "invalid_attributes", attribute_error)
            )
        images: list[str] = []
        if values.get("image_references"):
            try:
                images = [
                    safe_filename(value.strip())
                    for value in values["image_references"].split("|")
                    if value.strip()
                ]
            except ValueError:
                errors.append(
                    ParsedError(
                        number,
                        "image_references",
                        "invalid_image_reference",
                        "Image references must be plain filenames.",
                    )
                )
        metadata: dict[str, JsonValue] = {}
        if values.get("metadata"):
            try:
                decoded = json.loads(values["metadata"])
                if not isinstance(decoded, dict):
                    raise ValueError
                metadata = decoded
            except (json.JSONDecodeError, ValueError):
                errors.append(
                    ParsedError(
                        number,
                        "metadata",
                        "invalid_metadata",
                        "Metadata must be a JSON object.",
                    )
                )
        if errors:
            return None, errors
        json_attributes = cast(dict[str, JsonValue], attributes)
        json_images = cast(list[JsonValue], images)
        return {
            "operation": operation,
            "product_sku": (values["product_sku"]).upper(),
            "product_name": values["product_name"],
            "description": values.get("description") or None,
            "brand": values.get("brand") or None,
            "category": values["category"],
            "collection": values.get("collection") or None,
            "variant_sku": (values["variant_sku"]).upper(),
            "attributes": json_attributes,
            "price": format(price, "f"),
            "currency": currency,
            "quantity_on_hand": quantity,
            "image_references": json_images,
            "metadata": metadata,
        }, []


def safe_filename(value: str) -> str:
    normalized = value.strip()
    if (
        not normalized
        or len(normalized) > 255
        or "/" in normalized
        or "\\" in normalized
        or normalized in {".", ".."}
        or PurePosixPath(normalized).name != normalized
    ):
        raise ValueError("Filename is unsafe.")
    return normalized


def _header(value: object) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Spreadsheet headers must contain text.")
    return re.sub(r"[^a-z0-9]+", "_", value.strip().casefold()).strip("_")


def _cell(value: object) -> str:
    if value is None:
        return ""
    result = str(value).strip()
    if result.startswith(("=", "+", "@")) or (
        result.startswith("-") and not _numeric(result)
    ):
        raise ValueError("Formula-like spreadsheet values are not accepted.")
    return result


def _numeric(value: str) -> bool:
    try:
        Decimal(value)
    except InvalidOperation:
        return False
    return True


def _attributes(values: dict[str, str]) -> tuple[dict[str, str], str | None]:
    result: dict[str, str] = {}
    raw = values.get("attributes", "")
    if raw:
        for item in raw.split(";"):
            if "=" not in item:
                return (
                    {},
                    "Attributes must use key=value pairs separated by semicolons.",
                )
            key, value = item.split("=", 1)
            slug = re.sub(r"[^a-z0-9]+", "-", key.casefold()).strip("-")
            if not slug or not value.strip() or slug in result:
                return {}, "Attribute names and values must be non-empty and unique."
            result[slug] = value.strip()
    for key in ("size", "color"):
        if values.get(key):
            if key in result and result[key].casefold() != values[key].casefold():
                return {}, f"{key} conflicts with the attributes column."
            result[key] = values[key]
    if not result:
        return {}, "At least one normalized variant attribute is required."
    return dict(sorted(result.items())), None


def _optional_text(value: JsonValue | None) -> str | None:
    return str(value) if value is not None else None
